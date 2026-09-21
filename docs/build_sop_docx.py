from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = DOCS / "KC无限画布小白入门操作SOP.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 96, 110)
BLACK = RGBColor(0, 0, 0)


def set_run_font(run, name="Microsoft YaHei", size=None, color=None, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_run_font(r, size=9.2, color=BLACK, bold=bold)
    if fill:
        set_cell_shading(cell, fill)


def add_title(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("KC 无限画布小白入门操作 SOP")
    set_run_font(r, size=24, color=BLACK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("从进入系统、选择项目到新建画布、节点创作、保存与备份")
    set_run_font(r, size=11, color=MUTED)

    meta = doc.add_table(rows=4, cols=2)
    meta.autofit = False
    widths = [Inches(1.25), Inches(5.0)]
    rows = [
        ("适用对象", "第一次使用 KC 无限画布的业务人员、分镜/视频生成专员、项目负责人"),
        ("使用目标", "完成“选项目 -> 建画布 -> 添加节点 -> 上传/引用素材 -> 连线生成 -> 保存结果 -> 备份导出”的完整闭环"),
        ("系统定位", "内部影视与视觉创意业务的 AI 创作画布 MVP"),
        ("文档版本", "v1.0，基于当前 tapnow-base 原型与 PRD 截图整理"),
    ]
    for row, (label, value) in zip(meta.rows, rows):
        for idx, width in enumerate(widths):
            row.cells[idx].width = width
        set_cell_text(row.cells[0], label, bold=True, fill="E8EEF5")
        set_cell_text(row.cells[1], value)


def h1(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, size=16, color=BLUE, bold=True)


def h2(doc, text):
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_run_font(r, size=13, color=BLUE, bold=True)


def para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.2
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, size=10.5, color=BLACK, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r, size=10.5, color=BLACK)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.5, color=BLACK)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run_font(r, size=10.2, color=BLACK)


def number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run_font(r, size=10.2, color=BLACK)


def add_image(doc, rel_path, caption):
    image = DOCS / rel_path
    if not image.exists():
        para(doc, f"[缺少图片：{rel_path}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(image), width=Inches(6.25))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    set_run_font(r, size=8.8, color=MUTED)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.autofit = False
    for i, header in enumerate(headers):
        t.rows[0].cells[i].width = Inches(widths[i])
        set_cell_text(t.rows[0].cells[i], header, bold=True, fill="E8EEF5")
    for row_data in rows:
        row = t.add_row()
        for i, value in enumerate(row_data):
            row.cells[i].width = Inches(widths[i])
            set_cell_text(row.cells[i], value)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_flow_table(doc):
    table(
        doc,
        ["阶段", "用户动作", "系统结果"],
        [
            ("进入", "输入访问密码并点击进入画布", "进入项目管理页"),
            ("选项目", "筛选/搜索并点击项目卡片", "打开该项目的画布空间"),
            ("建画布", "打开子画布下拉并新建子画布", "获得独立创作空间"),
            ("创作", "添加文本/生图/生视频节点，上传素材并连线", "形成可追溯创作链路"),
            ("沉淀", "保存满意结果并导出备份", "结果进入项目素材/资产库，画布可恢复"),
        ],
        [0.85, 2.65, 2.8],
    )


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)

    add_title(doc)
    h1(doc, "0. 完整操作链路")
    para(doc, "新手只需要记住一个顺序：先进入系统，再选项目，再建子画布，最后围绕节点完成上传、连线、生成、保存和备份。")
    add_flow_table(doc)

    h1(doc, "1. 进入系统并登录")
    add_image(doc, Path("sop-screenshots/01-login.png"), "图 1：登录页")
    for item in ["打开 KC 画布系统地址。", "在“访问密码”输入框输入当前访问密码。", "点击“进入画布”。"]:
        number(doc, item)
    bullet(doc, "输入密码前按钮可能为灰色，输入后会变为可点击。")
    bullet(doc, "当前原型登录成功后会进入“KC 无限画布项目管理”页面。")

    h1(doc, "2. 在项目管理页选择项目")
    add_image(doc, Path("sop-screenshots/02-project-management-current.png"), "图 2：项目管理页")
    para(doc, "进入系统后，先确认项目归属。不要直接开始画图，避免素材、积分和备份归到错误项目。")
    for item in ["使用“全部项目组”“全部项目类型”筛选项目。", "在搜索框输入项目名称、画布名称或导演组关键词。", "点击目标项目卡片进入画布空间。"]:
        number(doc, item)

    h1(doc, "3. 首次进入项目后的确认弹窗")
    add_image(doc, Path("sop-screenshots/03-after-select-project.png"), "图 3：首次欢迎弹窗")
    para(doc, "首次进入画布会看到“欢迎使用 KC画布 MVP”的说明。阅读后点击“开始试用”，即可进入实际画布。")
    bullet(doc, "该弹窗说明当前试用范围和默认模型。")
    bullet(doc, "如果演示环境没有真实 API Key，部分生成可能走模拟结果。")

    h1(doc, "4. 认识画布主界面")
    add_image(doc, Path("sop-screenshots/06-canvas-main-current.png"), "图 4：画布主界面")
    table(
        doc,
        ["区域", "位置", "用途"],
        [
            ("返回项目列表", "左上角返回箭头", "回到项目管理页，切换其他项目"),
            ("项目与画布名", "左上角标题区", "确认当前正在操作的项目"),
            ("子画布选择", "标题下方“主画布”下拉", "切换或新建项目下的子画布"),
            ("左侧工具栏", "左侧竖排图标", "添加节点、生成历史、资产库、导入素材"),
            ("画布操作区", "中央网格区域", "摆放节点、拖拽、连线和生成"),
        ],
        [1.4, 2.0, 2.9],
    )

    h1(doc, "5. 新建或切换子画布")
    add_image(doc, Path("sop-screenshots/07-subcanvas-dropdown-current.png"), "图 5：子画布下拉")
    for item in ["点击左上角当前子画布名称，例如“主画布”。", "查看已有子画布并点击目标子画布切换。", "如需新空间，点击“新建子画布”。"]:
        number(doc, item)
    add_image(doc, Path("sop-screenshots/08-new-subcanvas-modal-current.png"), "图 6：新建子画布弹窗")
    para(doc, "命名建议：按分镜段落、资产方向或任务阶段命名，例如“第 1 集 03 场”“角色参考探索”“多角度精修”。")

    h1(doc, "6. 添加第一个节点")
    add_image(doc, Path("prd-screenshots/feature-list/04-add-node-panel.png"), "图 7：添加节点面板")
    table(
        doc,
        ["节点", "适合做什么", "新手场景"],
        [
            ("文本", "写剧本、提示词、创作说明", "先整理分镜描述或角色设定"),
            ("生图", "文本生图、参考图再生成", "生成角色、场景、道具图"),
            ("生视频", "文本/图片生成视频", "把分镜描述或参考图转成视频片段"),
            ("音频", "文本转语音或上传音频", "做旁白、语音参考或音频素材"),
        ],
        [1.0, 2.55, 2.75],
    )
    para(doc, "第一次使用建议先加一个“文本”节点，再加“生图”或“生视频”节点，用文本作为上游参考。")

    h1(doc, "7. 上传图片、视频或文本素材")
    add_image(doc, Path("prd-screenshots/feature-list/05-canvas-right-click-upload.png"), "图 8：右键上传菜单")
    for item in ["在画布空白处右键。", "选择上传入口。", "选择本地图片、视频或文本文件。", "系统按文件类型自动创建对应节点。"]:
        number(doc, item)
    bullet(doc, "支持图片、视频、.txt、.md、.markdown。")
    bullet(doc, "也可以拖入文件或复制图片后粘贴到画布。")

    h1(doc, "8. 使用文本、生图和生视频节点")
    h2(doc, "文本节点")
    add_image(doc, Path("prd-screenshots/feature-list/07-text-node-editor.png"), "图 9：文本节点编辑")
    para(doc, "文本节点适合承载剧本段落、创作说明、角色设定和提示词。提示词建议按“主体 -> 风格 -> 镜头 -> 约束”写。")
    h2(doc, "生图节点")
    add_image(doc, Path("prd-screenshots/feature-list/06-image-node-empty-params.png"), "图 10：生图节点参数")
    para(doc, "生图节点用于创建图片，也可以接收上游文本或图片作为参考。生成后可下载、保存、裁剪、多角度控制或继续连接到下游节点。")
    h2(doc, "生视频节点")
    add_image(doc, Path("prd-screenshots/feature-list/08-video-node-empty-params.png"), "图 11：生视频节点参数")
    para(doc, "生视频节点用于把文字描述、图片参考或多媒体素材转换为视频结果。若提示需要连接素材，先把上游图片或文本连到视频节点。")

    h1(doc, "9. 节点连线：把创作链路串起来")
    para(doc, "连线表示“谁作为谁的参考”。从上游参考节点拖到下游生成节点，形成可追溯的创作链路。")
    table(
        doc,
        ["典型链路", "用途"],
        [
            ("文本节点 -> 生图节点", "用剧本/提示词生成图片"),
            ("图片节点 -> 生图节点", "用参考图继续生成变体"),
            ("图片节点 -> 生视频节点", "用参考图生成视频"),
            ("文本节点 -> 生视频节点", "用分镜描述生成视频"),
        ],
        [2.3, 4.0],
    )

    h1(doc, "10. 复用资产、裁剪和多角度精修")
    h2(doc, "项目资产库")
    add_image(doc, Path("prd-screenshots/feature-list/09-asset-library-project.png"), "图 12：项目资产库")
    para(doc, "点击左侧“项目资产库”，可从项目库或公共库里添加角色、场景、道具。资产进入画布后会成为可继续编辑的标准图片节点。")
    h2(doc, "图片裁剪")
    add_image(doc, Path("prd-screenshots/feature-list/14-crop-editor.png"), "图 13：图片裁剪")
    para(doc, "选中图片节点后选择“图片裁剪”，设置画幅和裁剪区域，再点击“生成裁剪节点”。系统会保留来源节点并生成新的裁剪节点。")
    h2(doc, "多角度控制")
    add_image(doc, Path("prd-screenshots/feature-list/13-multi-angle-editor.png"), "图 14：多角度编辑器")
    para(doc, "多角度控制适合把角色图、场景图或道具图转成不同视角。可设置水平环绕、垂直俯仰、景别、输出画幅和补充提示词。")

    h1(doc, "11. 保存结果与导出备份")
    add_image(doc, Path("prd-screenshots/feature-list/15-save-result-modal.png"), "图 15：保存结果弹窗")
    table(
        doc,
        ["保存目标", "适合内容"],
        [
            ("项目素材", "临时结果、分镜过程图、项目内复用素材"),
            ("新建资产", "可复用角色、场景、道具"),
            ("更新资产", "已有资产的优化版本"),
        ],
        [1.6, 4.7],
    )
    add_image(doc, Path("prd-screenshots/15-backup-modal.png"), "图 16：导入/导出备份")
    para(doc, "重要项目建议每天结束前导出一次画布备份；大量生成或大改节点结构前，也先导出备份。")

    h1(doc, "12. 画布导航与整理")
    add_image(doc, Path("prd-screenshots/16-canvas-navigation-minimap.png"), "图 17：小地图与导航")
    bullet(doc, "用缩放按钮或滚轮调整视图。")
    bullet(doc, "用小地图快速定位内容区域。")
    bullet(doc, "节点从左到右排列：左边放参考，右边放生成结果。")
    bullet(doc, "不同探索方向建议拆成不同子画布。")

    h1(doc, "13. 新手练习任务")
    for item in [
        "登录系统。",
        "在项目管理页选择一个测试项目。",
        "点击“开始试用”。",
        "打开“主画布”下拉，新建“新手练习”子画布。",
        "添加文本节点，输入一段分镜描述。",
        "添加生图节点，把文本节点连接到生图节点。",
        "生成图片后做一次裁剪。",
        "对原图或裁剪图做一次多角度控制。",
        "把满意结果保存到项目素材。",
        "导出当前画布备份。",
    ]:
        number(doc, item)

    h1(doc, "14. 常见问题速查")
    table(
        doc,
        ["问题", "可能原因", "处理方式"],
        [
            ("登录后看不到画布", "还停留在项目管理页", "先选择项目卡片"),
            ("进入项目后被弹窗挡住", "首次欢迎弹窗未关闭", "点击“开始试用”"),
            ("找不到新建画布入口", "没打开子画布下拉", "点击“主画布”或当前子画布名称"),
            ("上传后节点类型不对", "文件类型和预期不一致", "检查文件是图片、视频还是文本"),
            ("节点连不上", "方向反了或类型不匹配", "从参考素材拖向生成节点"),
            ("结果找不到", "只在节点里生成，未保存", "保存到项目素材或资产库"),
            ("换电脑后没有画布", "本地数据未迁移", "使用画布备份导出/导入"),
        ],
        [1.55, 2.0, 2.75],
    )

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
