from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE_DIR = Path("docs/director-desk-prd")
SCREENSHOT_DIR = BASE_DIR / "screenshots"
OUT_DOCX = BASE_DIR / "director-desk-prd.docx"


ROWS = [
    {
        "level": "1\n1.1",
        "feature": "画布入口\n右键新建",
        "operation": "在画布空白处右键，菜单里点「3D导演台」。",
        "io": "输入：当前画布位置。\n输出：菜单项被选中。\n边界：点在节点上时不触发空白菜单；重复点击只新增新节点。",
        "shot": "01-canvas-context-menu.png",
    },
    {
        "level": "1\n1.2",
        "feature": "画布入口\n导演台节点",
        "operation": "点击菜单后，在右键位置旁边生成一个「3D导演台」节点，不自动连线。",
        "io": "输入：节点位置、instanceId。\n输出：导演台节点。\n边界：没有全景图也能打开；节点可移动、可删除。",
        "shot": "02-director-node.png",
    },
    {
        "level": "1\n1.3",
        "feature": "画布入口\n进入导演台",
        "operation": "点击导演台节点里的「进入」，打开一个单独的导演台空间。",
        "io": "输入：instanceId、主题、可选全景图。\n输出：导演台 iframe 页面。\n边界：加载失败时显示失败提示，可关闭返回画布。",
        "shot": "03-director-space-full.png",
    },
    {
        "level": "2\n2.1",
        "feature": "导演台空间\n基础界面",
        "operation": "进入后展示左侧对象列表、中间 3D 画面、右侧设置面板、底部工具栏。",
        "io": "输入：导演台实例数据。\n输出：可编辑的 3D 编排空间。\n边界：新实例默认有 1 个角色和 1 个机位。",
        "shot": "03-director-space-full.png",
    },
    {
        "level": "2\n2.2",
        "feature": "导演台空间\n场景设置",
        "operation": "点击空白场景或右侧「3D场景」，调整场景缩放、平移、旋转、天空色、全球球、地面等。",
        "io": "输入：数字、滑杆、开关、全景图。\n输出：场景视觉变化。\n边界：没有全景图时显示未连接；非法图片不导入。",
        "shot": "04-scene-panel.png",
    },
    {
        "level": "2\n2.3",
        "feature": "导演台空间\n对象树",
        "operation": "左侧按「角色」「摄像机」分组显示对象，可搜索、显示/隐藏、锁定/解锁。",
        "io": "输入：搜索词、可见开关、锁定开关。\n输出：对象列表和画面同步变化。\n边界：锁定后不允许误拖动；搜索为空恢复全部。",
        "shot": "03-director-space-full.png",
    },
    {
        "level": "2\n2.4",
        "feature": "导演台空间\n添加角色",
        "operation": "点底部「添加角色」，选择男性素体、女性素体、少年、儿童、群众、几何模型等。",
        "io": "输入：角色类型。\n输出：新角色或群众对象。\n边界：连续添加自动命名；群众对象按组合整体创建。",
        "shot": "05-add-character-menu.png",
    },
    {
        "level": "2\n2.5",
        "feature": "导演台空间\n模型库",
        "operation": "点底部「模型库」，从便利生活、居家生活、户外出行、工具配件、我的模型里选道具。",
        "io": "输入：模型分类和模型项。\n输出：道具进入 3D 场景。\n边界：模型为空时保留空状态；加载失败不影响已有场景。",
        "shot": "06-model-library.png",
    },
    {
        "level": "2\n2.6",
        "feature": "导演台空间\n角色属性",
        "operation": "点左侧角色，右侧可改名称、位置、旋转、缩放、颜色。",
        "io": "输入：名称、数字、颜色。\n输出：角色位置和外观变化。\n边界：数值为空不保存；锁定对象不可拖动。",
        "shot": "07-character-properties.png",
    },
    {
        "level": "2\n2.7",
        "feature": "导演台空间\n角色姿势",
        "operation": "选中角色后点「姿势」，选择站立、T型、行走、跑步等预设，也可用滑杆微调身体。",
        "io": "输入：姿势预设、身体滑杆。\n输出：角色动作姿态变化。\n边界：切换预设会覆盖当前姿势；微调只影响当前角色。",
        "shot": "08-pose-panel.png",
    },
    {
        "level": "2\n2.8",
        "feature": "导演台空间\n机位编辑",
        "operation": "点左侧机位，切换到「机位视角」，调整机位位置、注视目标、视野角度。",
        "io": "输入：机位、坐标、FOV。\n输出：镜头构图变化。\n边界：没有选中机位时不能截图；切换机位不丢失已有角色。",
        "shot": "09-camera-panel.png",
    },
    {
        "level": "2\n2.9",
        "feature": "导演台空间\n机位截图",
        "operation": "在机位面板点「摄像机截图」，生成当前机位预览，可清空或发送到画布。",
        "io": "输入：当前机位画面。\n输出：截图列表。\n边界：截图失败时保留场景；没有截图时「发送到画布」不可用。",
        "shot": "10-camera-screenshot-list.png",
    },
    {
        "level": "2\n2.10",
        "feature": "导演台空间\n视图工具",
        "operation": "用底部工具切换移动、旋转、添加图片、添加几何体、添加摄像机、网格、全屏等。",
        "io": "输入：工具按钮。\n输出：当前编辑模式变化。\n边界：正在输入数字时不触发画布快捷键。",
        "shot": "03-director-space-full.png",
    },
    {
        "level": "3\n3.1",
        "feature": "回到画布\n发送截图",
        "operation": "在导演台截图列表里点「发送到画布」。",
        "io": "输入：截图 dataURL。\n输出：发送给父级画布。\n边界：dataURL 为空或不是图片时丢弃；来源不匹配时拒收。",
        "shot": "10-camera-screenshot-list.png",
    },
    {
        "level": "3\n3.2",
        "feature": "回到画布\n图片节点",
        "operation": "画布收到截图后，在导演台节点旁边生成图片节点，不自动连线。",
        "io": "输入：dataURL、文件名、原图尺寸。\n输出：图片节点。\n边界：多张截图按顺序排列；图片尺寸按比例适配节点。",
        "shot": "11-return-image-node.png",
    },
    {
        "level": "4\n4.1",
        "feature": "复刻方案\n源码交付",
        "operation": "把开源导演台仓库作为独立 micro app 交给 Vue 画布团队，不直接混进画布主业务。",
        "io": "输入：原仓库源码、构建说明。\n输出：可单独运行、可单独构建的导演台子应用。\n边界：保留开源协议；不要改动核心 3D 能力再做迁移。",
        "shot": "03-director-space-full.png",
    },
    {
        "level": "4\n4.2",
        "feature": "复刻方案\n嵌入协议",
        "operation": "Vue 画布创建导演台节点，节点打开 iframe；父子页面用 postMessage 传数据和收截图。",
        "io": "输入：instanceId、theme、panorama。\n输出：captures 数组。\n边界：只接受同源或白名单来源；关闭 iframe 不删除节点。",
        "shot": "02-director-node.png",
    },
    {
        "level": "4\n4.3",
        "feature": "复刻方案\n落地验证",
        "operation": "先跑通「右键新建-进入导演台-截图-生成图片节点」最短链路，再补全道具、姿势、全景图。",
        "io": "输入：测试画布、测试全景图。\n输出：可验收链路。\n边界：不自动连线；截图分辨率不压缩；失败时不影响原画布功能。",
        "shot": "11-return-image-node.png",
    },
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths_in):
    table.allow_autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for idx, width in enumerate(widths_in):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def format_run(run, size=8.5, bold=False, color="000000"):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_text(paragraph, text: str, size=8.5, bold=False):
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.08
    for i, line in enumerate(text.split("\n")):
        if i:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        format_run(run, size=size, bold=bold)


def fill_text_cell(cell, text: str, size=8.5, bold=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    cell.text = ""
    set_paragraph_text(cell.paragraphs[0], text, size=size, bold=bold)


def fill_image_cell(cell, image_name: str):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    cell.text = ""
    image_path = SCREENSHOT_DIR / image_name
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(2.35))
    cap = cell.add_paragraph(image_name)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(0)
    format_run(cap.runs[0], size=7, color="555555")


def add_section_table(doc: Document, title: str, rows: list[dict]) -> None:
    heading = doc.add_heading(title, level=1)
    for run in heading.runs:
        format_run(run, size=16, bold=True, color="2E74B5")

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    widths = [0.55, 1.25, 2.05, 2.65, 2.55]
    headers = ["层级", "功能", "怎么操作", "输入 / 输出 / 边界", "原型截图"]
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, "E8EEF5")
        fill_text_cell(cell, header, size=9, bold=True)

    for item in rows:
        cells = table.add_row().cells
        fill_text_cell(cells[0], item["level"], size=8, bold=True)
        fill_text_cell(cells[1], item["feature"], size=8.5, bold=True)
        fill_text_cell(cells[2], item["operation"], size=8.5)
        fill_text_cell(cells[3], item["io"], size=8.2)
        fill_image_cell(cells[4], item["shot"])

    set_table_widths(table, widths)
    doc.add_paragraph()


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("3D导演台接入画布 - 功能需求方案")
    format_run(run, size=24, bold=True, color="000000")

    subtitle = doc.add_paragraph()
    set_paragraph_text(
        subtitle,
        "目标：把 3D 导演台作为画布里的一个功能点。用户在画布右键创建导演台节点，进入导演台完成 3D 编排，截图后回到画布生成图片节点。图片节点只放在旁边，不自动连线。",
        size=10,
    )

    grouped = {
        "一、画布入口": [row for row in ROWS if row["level"].startswith("1")],
        "二、导演台空间": [row for row in ROWS if row["level"].startswith("2")],
        "三、截图回到画布": [row for row in ROWS if row["level"].startswith("3")],
        "四、复刻到 Vue 画布": [row for row in ROWS if row["level"].startswith("4")],
    }
    for title, rows in grouped.items():
        add_section_table(doc, title, rows)

    doc.save(OUT_DOCX)
    print(OUT_DOCX.resolve())


if __name__ == "__main__":
    build()
